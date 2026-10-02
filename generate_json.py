import openpyxl, json

wb_orig = openpyxl.load_workbook(r'C:\Users\marum\OneDrive\動画\デスクトップ\10月マスタミス\shop_masters_202610020753_result1001.xlsx', data_only=True)
sheet_orig = wb_orig['店舗マスタ']

headers = [sheet_orig.cell(row=1, column=c).value for c in range(1, sheet_orig.max_column + 1)]

staff_list = ['下村\u3000公人', '吉田\u3000明香', '笹原\u3000彩佳', '山崎\u3000紗季', '小川\u3000菜緒子', '小野\u3000あかね', '秦\u3000貴凡', '垂水\u3000翔平', '清原\u3000由梨', '浅田\u3000鈴美', '大内\u3000真奈美', '本田\u3000宏一', '櫻井\u3000蓮']

staff_data = {}
for s in staff_list:
    staff_data[s] = {'a_count': 0, 'non_a_red_total': 0, 'shops': []}

col_red_ranking = {}

for r in range(2, sheet_orig.max_row + 1):
    shop = sheet_orig.cell(row=r, column=1).value
    staff = sheet_orig.cell(row=r, column=4).value or '未設定'
    sv = sheet_orig.cell(row=r, column=5).value or ''
    
    red_cols = []
    is_a_red = False
    
    for c in range(1, sheet_orig.max_column + 1):
        cell = sheet_orig.cell(row=r, column=c)
        fill = cell.fill
        if fill and fill.fill_type and fill.fill_type != 'none':
            rgb = getattr(fill.fgColor, 'rgb', None) or getattr(fill.start_color, 'rgb', None)
            if rgb == 'FFFF0000':
                col_name = headers[c-1] or f'Col{c}'
                red_cols.append(col_name)
                if c == 1:
                    is_a_red = True
                else:
                    col_red_ranking[col_name] = col_red_ranking.get(col_name, 0) + 1
                    
    non_a_red = [col for col in red_cols if col != '店舗']
    
    if staff not in staff_data:
        staff_data[staff] = {'a_count': 0, 'non_a_red_total': 0, 'shops': []}
        
    if is_a_red:
        staff_data[staff]['a_count'] += 1
    staff_data[staff]['non_a_red_total'] += len(non_a_red)
    
    staff_data[staff]['shops'].append({
        'name': shop,
        'sv': sv,
        'date': '10/01',
        'is_a_red': is_a_red,
        'non_a_count': len(non_a_red),
        'non_a_items': non_a_red
    })

# ミス率（%）の計算
for s, info in staff_data.items():
    total_shops = len(info['shops'])
    info['miss_rate'] = round((info['a_count'] / total_shops * 100), 1) if total_shops > 0 else 0.0

with open('portal_json_data.json', 'w', encoding='utf-8') as f:
    json.dump({'staff_data': staff_data, 'col_ranking': col_red_ranking}, f, ensure_ascii=False, indent=2)

print('Saved exact portal_json_data.json with miss_rate successfully')
