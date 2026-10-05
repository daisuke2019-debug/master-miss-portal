# -*- coding: utf-8 -*-
import openpyxl, json

# 1. Update portal_json_data.json with merged 10/01 + 10/02
import openpyxl, os

def parse_file(excel_path, date_str):
    tmp_path = 'tmp_' + os.path.basename(excel_path)
    with open(excel_path, 'rb') as f_in:
        with open(tmp_path, 'wb') as f_out:
            f_out.write(f_in.read())
            
    wb = openpyxl.load_workbook(tmp_path, data_only=True)
    sheet = wb['店舗マスタ'] if '店舗マスタ' in wb.sheetnames else wb.sheetnames[0]
    headers = [sheet.cell(row=1, column=c).value for c in range(1, sheet.max_column + 1)]
    
    file_shops = []
    
    for r in range(2, sheet.max_row + 1):
        shop = sheet.cell(row=r, column=1).value
        if not shop or shop == '店舗': continue
        
        staff = sheet.cell(row=r, column=4).value or '未設定'
        sv = sheet.cell(row=r, column=5).value or ''
        
        red_cols = []
        is_a_red = False
        
        for c in range(1, sheet.max_column + 1):
            cell = sheet.cell(row=r, column=c)
            fill = cell.fill
            if fill and fill.fill_type and fill.fill_type != 'none':
                fg = getattr(fill.fgColor, 'rgb', None) or getattr(fill.start_color, 'rgb', None)
                if fg:
                    fg_str = str(fg).upper()
                    if 'FF0000' in fg_str or fg_str == 'RED':
                        col_name = headers[c-1] or f'Col{c}'
                        red_cols.append(col_name)
                        if c == 1:
                            is_a_red = True
                            
        non_a_red = [col for col in red_cols if col != headers[0]]
        
        file_shops.append({
            'shop': shop,
            'staff': staff,
            'sv': sv,
            'date': date_str,
            'is_a_red': is_a_red,
            'non_a_count': len(non_a_red),
            'non_a_items': non_a_red
        })
        
    if os.path.exists(tmp_path): os.remove(tmp_path)
    return file_shops

f1 = r'C:\Users\marum\OneDrive\デスクトップ\10月マスタミス\shop_masters_202610020753_result1001.xlsx'
f2 = r'C:\Users\marum\OneDrive\デスクトップ\10月マスタミス\shop_masters_202610042003_result1002.xlsx'

shops_1001 = parse_file(f1, '10/01')
shops_1002 = parse_file(f2, '10/02')

merged_staff_data = {}
merged_col_ranking = {}

all_shops = shops_1001 + shops_1002

for item in all_shops:
    staff = item['staff']
    if staff not in merged_staff_data:
        merged_staff_data[staff] = {'a_count': 0, 'non_a_red_total': 0, 'shops': []}
        
    if item['is_a_red']:
        merged_staff_data[staff]['a_count'] += 1
    merged_staff_data[staff]['non_a_red_total'] += item['non_a_count']
    
    for col_name in item['non_a_items']:
        merged_col_ranking[col_name] = merged_col_ranking.get(col_name, 0) + 1
        
    merged_staff_data[staff]['shops'].append({
        'name': item['shop'],
        'sv': item['sv'],
        'date': item['date'],
        'is_a_red': item['is_a_red'],
        'non_a_count': item['non_a_count'],
        'non_a_items': item['non_a_items']
    })

for s, info in merged_staff_data.items():
    tot = len(info['shops'])
    info['miss_rate'] = round((info['a_count'] / tot * 100), 1) if tot > 0 else 0.0

with open('portal_json_data.json', 'w', encoding='utf-8') as f:
    json.dump({'staff_data': merged_staff_data, 'col_ranking': merged_col_ranking}, f, ensure_ascii=False, indent=2)

print('Updated portal_json_data.json with 10/01 + 10/02 merged data.')
