# -*- coding: utf-8 -*-
"""10月マスタミスフォルダ内の *_resultMMDD.xlsx を全自動検出して portal_json_data.json を再生成"""
import openpyxl, json, os, re, sys, ctypes, ctypes.wintypes
sys.stdout.reconfigure(encoding='utf-8')

SRC_DIR = r'C:\Users\marum\OneDrive\デスクトップ\10月マスタミス'
WORK_DIR = os.path.dirname(os.path.abspath(__file__))

def safe_copy(src, dst):
    k = ctypes.windll.kernel32
    k.CreateFileW.restype = ctypes.wintypes.HANDLE
    h = k.CreateFileW(src, 0x80000000, 0x7, None, 3, 0x80, None)
    if h in (-1, 0xFFFFFFFFFFFFFFFF, None):
        raise PermissionError(src)
    o = k.CreateFileW(dst, 0x40000000, 0, None, 2, 0x80, None)
    buf = ctypes.create_string_buffer(65536)
    r = ctypes.wintypes.DWORD(); w = ctypes.wintypes.DWORD()
    while k.ReadFile(h, buf, 65536, ctypes.byref(r), None) and r.value:
        k.WriteFile(o, buf, r.value, ctypes.byref(w), None)
    k.CloseHandle(h); k.CloseHandle(o)

def parse_file(path, date_str):
    tmp = os.path.join(WORK_DIR, 'tmp_scan.xlsx')
    safe_copy(path, tmp)
    wb = openpyxl.load_workbook(tmp, data_only=True)
    sh = wb['店舗マスタ'] if '店舗マスタ' in wb.sheetnames else wb[wb.sheetnames[0]]
    headers = [str(sh.cell(row=1, column=c).value or f'Col{c}').strip() for c in range(1, sh.max_column + 1)]
    shops = []
    for r in range(2, sh.max_row + 1):
        v = sh.cell(row=r, column=1).value
        if not v or str(v).strip() == '店舗':
            continue
        staff = str(sh.cell(row=r, column=4).value or '未設定').strip()
        sv = str(sh.cell(row=r, column=5).value or '').strip()
        is_a_red, non_a = False, []
        for c in range(1, sh.max_column + 1):
            f = sh.cell(row=r, column=c).fill
            if f and f.fill_type and f.fill_type != 'none':
                fg = str(getattr(f.fgColor, 'rgb', '') or getattr(f.start_color, 'rgb', '') or '').upper()
                if 'FF0000' in fg or fg == 'RED':
                    if c == 1:
                        is_a_red = True
                    else:
                        non_a.append(headers[c - 1])
        shops.append({'name': str(v).strip(), 'staff': staff, 'sv': sv, 'date': date_str,
                      'is_a_red': is_a_red, 'non_a_count': len(non_a), 'non_a_items': non_a})
    wb.close()
    os.remove(tmp)
    return shops

files = {}
for fn in os.listdir(SRC_DIR):
    m = re.search(r'_result(\d{2})(\d{2})\.xlsx$', fn, re.I)
    if m:
        files[f'{m.group(1)}/{m.group(2)}'] = os.path.join(SRC_DIR, fn)

staff_data, col_ranking = {}, {}
for d in sorted(files):
    shops = parse_file(files[d], d)
    print(f'{d}: {len(shops)}店舗 ({os.path.basename(files[d])})')
    for it in shops:
        s = staff_data.setdefault(it['staff'], {'a_count': 0, 'non_a_red_total': 0, 'shops': []})
        s['a_count'] += it['is_a_red']
        s['non_a_red_total'] += it['non_a_count']
        for col in it['non_a_items']:
            col_ranking[col] = col_ranking.get(col, 0) + 1
        s['shops'].append(it)

for info in staff_data.values():
    t = len(info['shops'])
    info['miss_rate'] = round(info['a_count'] / t * 100, 1) if t else 0.0

with open(os.path.join(WORK_DIR, 'portal_json_data.json'), 'w', encoding='utf-8') as f:
    json.dump({'staff_data': staff_data, 'col_ranking': col_ranking,
               'available_dates': sorted(files)}, f, ensure_ascii=False, indent=2, default=str)
print('合計:', sum(len(v['shops']) for v in staff_data.values()), '店舗')
