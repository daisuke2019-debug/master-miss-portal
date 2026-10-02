import openpyxl

wb_orig = openpyxl.load_workbook(r'C:\Users\marum\OneDrive\動画\デスクトップ\10月マスタミス\shop_masters_202610020753_result1001.xlsx', data_only=True)
sheet_orig = wb_orig['店舗マスタ']

wb_master = openpyxl.load_workbook(r'C:\Users\marum\OneDrive\動画\デスクトップ\10月マスタミス\【マスタミス管理マスター】10月度.xlsx', data_only=True)
sheet_log = wb_master['🔴ピンポイント赤マーカー指摘ログ']

headers = [sheet_orig.cell(row=1, column=c).value for c in range(1, sheet_orig.max_column + 1)]

with open('compare_report.txt', 'w', encoding='utf-8') as out:
    out.write('=== 全40店舗の【純赤色#FFFF0000 セル数】と【管理マスター記載数】の完全対比 ===\n\n')
    out.write(f'{"店舗名":<15} | {"スタッフ":<10} | {"実ファイル全赤セル数":<18} | {"実ファイルA列外赤セル数":<20} | {"管理マスター記載数":<15}\n')
    out.write('-'*90 + '\n')

    for r in range(2, sheet_orig.max_row + 1):
        shop = sheet_orig.cell(row=r, column=1).value
        staff = sheet_orig.cell(row=r, column=4).value or ''
        
        all_red_cols = []
        for c in range(1, sheet_orig.max_column + 1):
            cell = sheet_orig.cell(row=r, column=c)
            fill = cell.fill
            if fill and fill.fill_type and fill.fill_type != 'none':
                rgb = getattr(fill.fgColor, 'rgb', None) or getattr(fill.start_color, 'rgb', None)
                if rgb == 'FFFF0000':
                    all_red_cols.append(headers[c-1])
                    
        non_a_red = [col for col in all_red_cols if col != '店舗']
        
        master_cnt = 'なし'
        for mr in range(4, sheet_log.max_row + 1):
            if sheet_log.cell(row=mr, column=2).value == shop:
                master_cnt = sheet_log.cell(row=mr, column=5).value
                break
                
        out.write(f'{str(shop):<15} | {str(staff):<10} | {len(all_red_cols):<18} | {len(non_a_red):<20} | {str(master_cnt):<15}\n')

print('Wrote compare_report.txt successfully')
